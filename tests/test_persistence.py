from pathlib import Path

from bc2800.domain.models import Exam
from bc2800.persistence.excel_writer import ExcelWriter
from bc2800.persistence.sqlite_repo import SqliteRepo
from bc2800.protocol.parser import parse_exam

from frames import live_cdh_dog, make_cdh, make_farm, make_goat


def test_sqlite_insert_and_dedup(tmp_path: Path) -> None:
    repo = SqliteRepo(tmp_path / "t.sqlite")
    exam = parse_exam(make_cdh())
    stored, inserted = repo.insert_exam(exam)
    assert inserted
    assert stored.db_id is not None
    again, inserted_again = repo.insert_exam(parse_exam(make_cdh()))
    assert not inserted_again
    assert again.db_id == stored.db_id
    assert repo.pending_excel_count() == 1
    repo.mark_excel_synced(stored.db_id)
    assert repo.pending_excel_count() == 0
    repo.close()


def test_sqlite_user_fields(tmp_path: Path) -> None:
    repo = SqliteRepo(tmp_path / "t.sqlite")
    exam, _ = repo.insert_exam(parse_exam(make_farm()))
    updated = repo.update_user_fields(exam.db_id, "Mimosa", "Ana", "OK")
    assert updated is not None
    assert updated.animal_name == "Mimosa"
    assert updated.tutor == "Ana"
    repo.close()


def test_excel_append_unified_columns(tmp_path: Path) -> None:
    path = tmp_path / "exames.xlsx"
    writer = ExcelWriter(path)
    dog = parse_exam(make_cdh())
    pig = parse_exam(make_farm())
    goat = parse_exam(make_goat())
    writer.append_exam(dog)
    writer.append_exam(pig)
    writer.append_exam(goat)
    from openpyxl import load_workbook

    book = load_workbook(path)
    sheet = book["Exames"]
    assert sheet.cell(1, 24).value == "Nome do animal"
    assert sheet.cell(1, 14).value == "HGB (g/dL)"
    assert sheet.cell(1, 15).value == "MCHC (g/dL)"
    assert sheet.cell(2, 3).value == "00000123"
    assert sheet.cell(2, 14).value == 14.2
    assert sheet.cell(2, 4).value == "Cão"
    assert sheet.cell(2, 7).value == 4.1
    assert sheet.cell(3, 4).value == "Suíno"
    assert sheet.cell(3, 7).value is None
    assert sheet.cell(4, 4).value == "Caprino"
    assert sheet.cell(4, 20).value is None


def test_reparse_failed_live_exam(tmp_path: Path) -> None:
    repo = SqliteRepo(tmp_path / "t.sqlite")
    raw = live_cdh_dog()
    stored, inserted = repo.insert_exam(
        Exam(raw_payload=raw, parse_ok=False, layout="unknown")
    )
    assert inserted
    assert not stored.parse_ok
    updated = repo.reparse_failed_exams()
    assert len(updated) == 1
    exam = updated[0]
    assert exam.parse_ok
    assert exam.layout == "cdh"
    assert exam.sample_id == "00002007"
    assert exam.wbc == 27.4
    assert exam.hgb == 13.5
    assert exam.plt == 430
    repo.close()


def test_insert_repairs_failed_duplicate(tmp_path: Path) -> None:
    repo = SqliteRepo(tmp_path / "t.sqlite")
    raw = live_cdh_dog()
    repo.insert_exam(Exam(raw_payload=raw, parse_ok=False, layout="unknown"))
    stored, inserted = repo.insert_exam(parse_exam(raw))
    assert inserted
    assert stored.parse_ok
    assert stored.hgb == 13.5
    repo.close()


def test_excel_fill_blank_row(tmp_path: Path) -> None:
    path = tmp_path / "exames.xlsx"
    writer = ExcelWriter(path)
    raw = live_cdh_dog()
    blank = Exam(raw_payload=raw, parse_ok=False, layout="unknown")
    writer.append_exam(blank)
    parsed = parse_exam(raw)
    parsed.received_at = blank.received_at
    assert writer.fill_exam_row(parsed)
    from openpyxl import load_workbook

    book = load_workbook(path)
    sheet = book["Exames"]
    assert sheet.cell(2, 3).value == "00002007"
    assert sheet.cell(2, 6).value == 27.4
    assert sheet.cell(2, 14).value == 13.5
    assert sheet.max_row == 2
