from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bc2800.domain.models import Exam
from bc2800.domain.species import format_exam_at, sample_mode_label, species_name

_LOG = logging.getLogger("bc2800.excel")

SHEET = "Exames"

HEADERS = [
    "Recebido em",
    "Data/hora do exame",
    "ID",
    "Espécie",
    "Modo da amostra",
    "WBC (10^9/L)",
    "Lymph# (10^9/L)",
    "Mid# (10^9/L)",
    "Gran# (10^9/L)",
    "Lymph% (%)",
    "Mid% (%)",
    "Gran% (%)",
    "RBC (10^12/L)",
    "HGB (g/dL)",
    "MCHC (g/dL)",
    "MCV (fL)",
    "MCH (pg)",
    "RDW (%)",
    "HCT (%)",
    "PLT (10^9/L)",
    "MPV (fL)",
    "PDW",
    "PCT (%)",
    "Nome do animal",
    "Tutor",
    "Observações",
]

COL_ID = 3
COL_EXAM_AT = 2
COL_RECEIVED = 1
COL_SPECIES = 4
COL_NAME = 24
COL_TUTOR = 25
COL_NOTES = 26


def _rename_unit_headers(sheet: Worksheet) -> bool:
    changed = False
    if sheet.cell(1, 14).value == "HGB (g/L)":
        sheet.cell(1, 14).value = "HGB (g/dL)"
        changed = True
    if sheet.cell(1, 15).value == "MCHC (g/L)":
        sheet.cell(1, 15).value = "MCHC (g/dL)"
        changed = True
    return changed


def _header_fill() -> PatternFill:
    return PatternFill("solid", fgColor="1F4E5F")


def _ensure_workbook(path: Path) -> Worksheet:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        book = load_workbook(path)
        if SHEET in book.sheetnames:
            sheet = book[SHEET]
            if _rename_unit_headers(sheet):
                book.save(path)
            return sheet
        sheet = book.create_sheet(SHEET)
        _write_headers(sheet)
        book.save(path)
        return sheet
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    _write_headers(sheet)
    book.save(path)
    return sheet


def _write_headers(sheet: Worksheet) -> None:
    fill = _header_fill()
    font = Font(color="FFFFFF", bold=True)
    for index, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, index, title)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = 16
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["X"].width = 22
    sheet.column_dimensions["Y"].width = 22
    sheet.column_dimensions["Z"].width = 32
    sheet.cell(1, 5).comment = Comment(
        "0 = Sangue total; 1 = Pré-diluído (valores típicos da BC-2800).",
        "BC2800 Receptor",
    )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"


def _cell(value: object) -> object:
    if value is None:
        return None
    return value


def exam_row(exam: Exam) -> list[object]:
    return [
        exam.received_at.strftime("%d/%m/%Y %H:%M:%S"),
        format_exam_at(exam.exam_at),
        exam.sample_id,
        species_name(exam.animal_type),
        sample_mode_label(exam.sample_mode),
        _cell(exam.wbc),
        _cell(exam.lymph_abs),
        _cell(exam.mid_abs),
        _cell(exam.gran_abs),
        _cell(exam.lymph_pct),
        _cell(exam.mid_pct),
        _cell(exam.gran_pct),
        _cell(exam.rbc),
        _cell(exam.hgb),
        _cell(exam.mchc),
        _cell(exam.mcv),
        _cell(exam.mch),
        _cell(exam.rdw),
        _cell(exam.hct),
        _cell(exam.plt),
        _cell(exam.mpv),
        _cell(exam.pdw),
        _cell(exam.pct),
        exam.animal_name,
        exam.tutor,
        exam.notes,
    ]


class ExcelLockedError(Exception):
    pass


class ExcelWriter:
    def __init__(self, path: Path) -> None:
        self.path = path

    def append_exam(self, exam: Exam) -> None:
        try:
            _ensure_workbook(self.path)
            book = load_workbook(self.path)
            sheet = book[SHEET] if SHEET in book.sheetnames else book.create_sheet(SHEET)
            if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
                _write_headers(sheet)
            _rename_unit_headers(sheet)
            row_values = exam_row(exam)
            sheet.append(row_values)
            row_idx = sheet.max_row
            id_cell = sheet.cell(row_idx, COL_ID)
            id_cell.number_format = "@"
            id_cell.value = exam.sample_id
            last = get_column_letter(len(HEADERS))
            sheet.auto_filter.ref = f"A1:{last}{sheet.max_row}"
            book.save(self.path)
        except PermissionError as exc:
            raise ExcelLockedError(str(self.path)) from exc

    def _find_row(self, sheet: Worksheet, exam: Exam) -> int | None:
        exam_label = format_exam_at(exam.exam_at)
        species = species_name(exam.animal_type)
        received = exam.received_at.strftime("%d/%m/%Y %H:%M:%S")
        received_match: int | None = None
        for row in range(2, sheet.max_row + 1):
            id_ok = str(sheet.cell(row, COL_ID).value or "") == (exam.sample_id or "")
            at_ok = str(sheet.cell(row, COL_EXAM_AT).value or "") == exam_label
            sp_ok = str(sheet.cell(row, COL_SPECIES).value or "") == species
            if exam.sample_id and id_ok and at_ok and sp_ok:
                return row
            if str(sheet.cell(row, COL_RECEIVED).value or "") == received:
                received_match = row
        return received_match

    def fill_exam_row(self, exam: Exam) -> bool:
        if not self.path.exists():
            return False
        try:
            book = load_workbook(self.path)
            if SHEET not in book.sheetnames:
                return False
            sheet = book[SHEET]
            _rename_unit_headers(sheet)
            target = self._find_row(sheet, exam)
            if target is None:
                book.save(self.path)
                return False
            for col, value in enumerate(exam_row(exam), start=1):
                sheet.cell(target, col, value)
            id_cell = sheet.cell(target, COL_ID)
            id_cell.number_format = "@"
            id_cell.value = exam.sample_id
            book.save(self.path)
            return True
        except PermissionError as exc:
            raise ExcelLockedError(str(self.path)) from exc

    def update_user_fields(self, exam: Exam) -> bool:
        if not self.path.exists() or exam.sample_id is None:
            return False
        try:
            book = load_workbook(self.path)
            sheet = book[SHEET]
            target = None
            exam_label = format_exam_at(exam.exam_at)
            species = species_name(exam.animal_type)
            for row in range(2, sheet.max_row + 1):
                if (
                    str(sheet.cell(row, COL_ID).value or "") == (exam.sample_id or "")
                    and str(sheet.cell(row, COL_EXAM_AT).value or "") == exam_label
                    and str(sheet.cell(row, COL_SPECIES).value or "") == species
                ):
                    target = row
                    break
            if target is None:
                return False
            sheet.cell(target, COL_NAME, exam.animal_name)
            sheet.cell(target, COL_TUTOR, exam.tutor)
            sheet.cell(target, COL_NOTES, exam.notes)
            book.save(self.path)
            return True
        except PermissionError:
            return False
